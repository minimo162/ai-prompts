"""Read-only Excel acceptance oracle. A logical PASS is not Copilot/PAD acceptance."""
import argparse
import hashlib
import json
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries, coordinate_to_tuple, get_column_letter

BASE = Path(__file__).resolve().parents[1] / 'catalog/acceptance/issue38'


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def rectangle(address):
    import re
    if not isinstance(address, str) or not re.fullmatch(r'[A-Z]+[1-9][0-9]*:[A-Z]+[1-9][0-9]*', address):
        raise ValueError('Invalid fixed rectangle')
    c1, r1, c2, r2 = range_boundaries(address)
    if not (1 <= c1 <= c2 <= 16384 and 1 <= r1 <= r2 <= 1048576):
        raise ValueError('Invalid rectangle bounds')
    if (c2-c1+1)*(r2-r1+1) > 10000:
        raise ValueError('Outside bounded fixture contract')
    return c1, r1, c2, r2


def typed(value):
    if value is None:
        return ['blank', None]
    if isinstance(value, bool):
        return ['boolean', value]
    if isinstance(value, (int, float)):
        return ['number', value]
    if isinstance(value, str):
        return ['text', value]
    raise ValueError(f'Unverified cell type: {type(value).__name__}')


def style(cell):
    # Compare semantic styles, never workbook-local style IDs.
    return [str(cell.font), str(cell.fill), str(cell.border), str(cell.alignment),
            cell.number_format, str(cell.protection)]


def cell_value(cell):
    return ['formula', cell.value] if cell.data_type == 'f' else typed(cell.value)


def case_spec(case_id):
    alias = 'EX02' if case_id == 'EX01' else case_id
    return next(c for c in json.loads((BASE/'spec.json').read_text(encoding='utf-8'))['cases'] if c['id'] == alias)


def check_frozen():
    frozen = json.loads((BASE/'freeze.json').read_text(encoding='utf-8'))
    for relative, expected in frozen['files'].items():
        if sha(BASE/relative) != expected:
            raise ValueError(f'Frozen input changed: {relative}')
    return frozen


def targets(spec):
    matrices = json.loads((BASE/'expected.json').read_text(encoding='utf-8'))['matrices'][spec['id']]
    result = {}
    for item, matrix in zip(spec['inputs'], matrices):
        c1, r1, c2, r2 = rectangle(item['range'])
        if len(matrix) != r2-r1+1 or any(len(row) != c2-c1+1 for row in matrix):
            raise ValueError('Expected matrix shape mismatch')
        r, c = coordinate_to_tuple(item['target_start'])
        if r+len(matrix)-1 > 1048576 or c+len(matrix[0])-1 > 16384:
            raise ValueError('Target out of bounds')
        for y, row in enumerate(matrix):
            for x, value in enumerate(row):
                key = (item['target_sheet'], f'{get_column_letter(c+x)}{r+y}')
                if key in result:
                    raise ValueError('Overlapping targets')
                result[key] = typed(value)
    return result


def verify_fixtures():
    counts = {}
    for case_id in ['EX02','EX03']:
        spec = case_spec(case_id)
        matrices = json.loads((BASE/'expected.json').read_text(encoding='utf-8'))['matrices'][case_id]
        for item, expected in zip(spec['inputs'], matrices):
            wb = load_workbook(BASE/item['file'], data_only=False)
            try:
                assert wb.active.title == 'Decoy', 'Wrong initial sheet'
                assert item['sheet'] in wb.sheetnames
                actual = [[typed(c.value) for c in row] for row in wb[item['sheet']][item['range']]]
                assert actual == [[typed(v) for v in row] for row in expected]
                assert all(c.data_type != 'f' for row in wb[item['sheet']][item['range']] for c in row)
            finally:
                wb.close()
        template = load_workbook(BASE/spec['template'], data_only=False)
        try:
            assert template.active.title == 'Decoy'
            for item in spec['inputs']:
                sheet = template[item['target_sheet']]
                assert sheet['A15'].value == 'KEEP OUTSIDE'
                assert sheet['B15'].value == '=2+3'
                assert sheet['B15'].number_format == '0.00'
        finally:
            template.close()
        counts[case_id] = len(targets(spec))
    return {'status':'PASS_FIXTURE_READBACK_ONLY','target_cells':counts,'pad_run':'NOT_RUN'}


def preflight(spec, output, work):
    check_frozen()
    originals = [BASE/item['file'] for item in spec['inputs']] + [BASE/spec['template']]
    output, work = Path(output), Path(work)
    if output.suffix.lower() != '.xlsx' or work.suffix.lower() != '.xlsx':
        raise ValueError('Only xlsx')
    all_paths = originals + [output, work]
    resolved = [str(p.resolve()).casefold() for p in all_paths]
    if len(set(resolved)) != len(resolved):
        raise ValueError('Input/template/work/output must differ')
    for i, a in enumerate(all_paths):
        for b in all_paths[i+1:]:
            if a.exists() and b.exists() and a.samefile(b):
                raise ValueError('Hardlink alias')
    if output.exists() or work.exists():
        raise ValueError('Existing work/output collision: stop, never overwrite')
    for item in spec['inputs']:
        rectangle(item['range'])
        wb = load_workbook(BASE/item['file'], read_only=True)
        try:
            if item['sheet'] not in wb.sheetnames:
                raise ValueError('Missing input sheet')
        finally:
            wb.close()
    template = load_workbook(BASE/spec['template'], read_only=True)
    try:
        for name, _ in targets(spec):
            if name not in template.sheetnames:
                raise ValueError('Missing target sheet')
    finally:
        template.close()
    return originals


def compare(case_id, output):
    check_frozen()
    spec = case_spec(case_id)
    template = load_workbook(BASE/spec['template'], data_only=False)
    actual = load_workbook(output, data_only=False)
    changes = targets(spec)
    failures, count = [], 0
    try:
        if actual.sheetnames != template.sheetnames:
            failures.append('sheet names/order changed')
        for ts in template:
            if ts.title not in actual.sheetnames:
                continue
            os = actual[ts.title]
            if str(ts.merged_cells) != str(os.merged_cells):
                failures.append(f'{ts.title}: merged cells changed')
            # Ignore selection/active tab: explicit sheet switching is expected.
            for coordinate in set(ts.row_dimensions) | set(os.row_dimensions):
                a, b = ts.row_dimensions[coordinate], os.row_dimensions[coordinate]
                if (a.height,a.hidden) != (b.height,b.hidden):
                    failures.append(f'{ts.title}: row dimension {coordinate}')
            for coordinate in set(ts.column_dimensions) | set(os.column_dimensions):
                a, b = ts.column_dimensions[coordinate], os.column_dimensions[coordinate]
                if (a.width,a.hidden,a.min,a.max) != (b.width,b.hidden,b.min,b.max):
                    failures.append(f'{ts.title}: column dimension {coordinate}')
            max_row,max_col=max(ts.max_row,os.max_row),max(ts.max_column,os.max_column)
            if max_row*max_col > 100000:
                failures.append('Unexpected sheet size'); continue
            for row in os.iter_rows(max_row=max_row,max_col=max_col):
                for cell in row:
                    before=ts[cell.coordinate]
                    wanted=changes.get((ts.title,cell.coordinate),cell_value(before))
                    if cell_value(cell) != wanted:
                        failures.append(f'{ts.title}!{cell.coordinate}: value/type/position')
                    if style(cell) != style(before):
                        failures.append(f'{ts.title}!{cell.coordinate}: style changed')
                    if (cell.comment.text if cell.comment else None) != (before.comment.text if before.comment else None):
                        failures.append(f'{ts.title}!{cell.coordinate}: comment changed')
                    count += 1
    finally:
        template.close(); actual.close()
    return {'status':'FAIL' if failures else 'PASS_LOGICAL_OUTPUT_ONLY','case':case_id,
            'checked_cells':count,'target_cells':len(changes),'failures':failures,
            'output_sha256':sha(output),'copilot_generation':'NOT_ASSERTED','pad_execution':'NOT_ASSERTED'}


if __name__ == '__main__':
    parser=argparse.ArgumentParser()
    parser.add_argument('--fixtures',action='store_true')
    parser.add_argument('--case',choices=['EX02','EX03'])
    parser.add_argument('--output',type=Path)
    args=parser.parse_args()
    result=verify_fixtures() if args.fixtures else compare(args.case,args.output)
    print(json.dumps(result,ensure_ascii=False,indent=2))
    raise SystemExit(1 if result['status']=='FAIL' else 0)
