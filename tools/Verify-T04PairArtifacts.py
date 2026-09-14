"""Read-only fixed T04 CSV/xlsx acceptance; never edits workbook bytes."""
import csv, datetime, hashlib, json, pathlib, sys
from openpyxl import load_workbook

def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()

def verify(root, directory, number):
    plan = json.loads((directory / 'plan.json').read_text(encoding='utf-8-sig'))
    observation = json.loads((directory / f'run{number}-observation.json').read_text(encoding='utf-8-sig'))
    assert number in (1, 2) and plan['max_run_starts'] == 2
    assert observation['verification_id'] == plan['verification_id']
    assert observation['run_number'] == number and observation['completion_observed']
    assert observation['running_observed'] and not observation['additional_run_requested']
    assert any(x.get('running') is True for x in observation['samples'])
    final = observation['samples'][-1]
    assert final['ready'] and not final['running'] and final['errors_known'] and final['errors'] == 0
    assert sha(root / plan['input']) == plan['input_sha256']
    expected = plan['expected_cells']
    assert expected == [['対象', 'A', 10], ['対象', 'C', 25], ['対象', 'D', 5]]
    artifacts = []
    observed = {}
    for name in plan['outputs']:
        p = directory / f'run{number}' / pathlib.Path(name).name
        gate = json.loads(p.with_name(p.name + '.gate.json').read_text(encoding='utf-8-sig'))
        assert gate['status'] == 'PASS' and gate['run_number'] == number
        assert gate['input_unchanged'] and gate['snapshot_sha256'] == sha(p)
        assert gate['snapshot_bytes'] == p.stat().st_size > 0
        if p.suffix == '.csv':
            with p.open(encoding='utf-8-sig', newline='') as f:
                cells = list(csv.reader(f))
            assert cells == [[str(x) for x in row] for row in expected], cells
        elif p.suffix == '.xlsx':
            wb = load_workbook(p, read_only=True, data_only=True)
            try:
                assert len(wb.worksheets) == 1
                ws = wb.worksheets[0]
                assert ws.max_row == 3 and ws.max_column == 3
                cells = [list(row) for row in ws.iter_rows(values_only=True)]
                assert cells == expected, cells
                observed['sheet'] = ws.title
            finally:
                wb.close()
        else:
            raise AssertionError('unexpected artifact type')
        observed[p.suffix] = cells
        artifacts.append({'path': p.relative_to(root).as_posix(), 'sha256': sha(p), 'bytes': p.stat().st_size})
    assert len(artifacts) == 2
    return {'status': 'PASS', 'verification_id': plan['verification_id'], 'run_number': number,
            'verified_at': datetime.datetime.now(datetime.timezone.utc).isoformat(),
            'expected': expected, 'observed': observed, 'artifacts': artifacts,
            'input_unchanged': True, 'gate': 'READY_FOR_NEXT_RUN' if number == 1 else 'PAIR_ARTIFACTS_COMPLETE'}

if __name__ == '__main__':
    root = pathlib.Path(__file__).resolve().parent.parent
    directory = (root / sys.argv[1]).resolve()
    number = int(sys.argv[2])
    output = directory / f'run{number}' / 'values.json'
    assert not output.exists(), 'Never overwrite an earlier value decision'
    result = verify(root, directory, number)
    with output.open('x', encoding='utf-8') as f:
        json.dump(result, f, ensure_ascii=False, indent=2)
        f.write('\n')
    print(f'PASS Run{number}: CSV and xlsx fixed 3x3 cells, hashes and input invariance')
